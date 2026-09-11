"""Reglas y bitácora del agente CFDI por factura. No accede al portal."""
from __future__ import annotations

import hashlib
import json
import re
import sqlite3
import unicodedata
from datetime import datetime
from pathlib import Path


class RevisionNecesaria(ValueError):
    pass


def normalizar(texto: str) -> str:
    texto = unicodedata.normalize("NFD", texto)
    return " ".join("".join(c for c in texto if not unicodedata.combining(c)).upper().split())


def documento(texto: str) -> tuple[str, int]:
    # Las capturas muestran 000DEM001001 en cobranza y DEM001001 en ventas.
    # Sólo se permite ese relleno; nunca coincidencias parciales ni quitar serie.
    match = re.fullmatch(r"0*([A-Z]+)[ -]?(\d+)", normalizar(texto))
    if not match:
        raise RevisionNecesaria(f"Formato de factura no reconocido: {texto!r}")
    return match[1], int(match[2])


def analizar_texto(texto: str, cliente: str, factura: str) -> dict:
    t = normalizar(texto)
    metodos = set()
    evidencias = []
    # Limitar al campo Método de pago: no confundir Forma de pago ni descripciones.
    for match in re.finditer(
        r"METODO\s+(?:DE\s+)?PAGO\s*[:\-]?\s*"
        r"((?:(?:PUE|PPD)\b\s*[-:]?\s*)?"
        r"PAGO\s+(?:EN\s+UNA\s+SOLA\s+EXHIBICION|EN\s+PARCIALIDADES\s+O\s+DIFERIDO)"
        r"|PUE\b|PPD\b)", t
    ):
        valor = match[1]
        if re.search(r"\bPUE\b|UNA SOLA EXHIBICION", valor):
            metodos.add("PUE")
        if re.search(r"\bPPD\b|PARCIALIDADES O DIFERIDO", valor):
            metodos.add("PPD")
        evidencias.append(match[0])
    if len(metodos) != 1:
        raise RevisionNecesaria("Método de pago ausente, ilegible o contradictorio en el PDF.")
    series = set(re.findall(r"\bSERIE\s*:\s*([A-Z]+)\b", t))
    folios = {int(x) for x in re.findall(r"\bFOLIO\s*:\s*(\d+)\b", t)}
    serie, folio = documento(factura)
    if series != {serie} or folios != {folio}:
        raise RevisionNecesaria("Serie/folio del PDF no coincide inequívocamente con la factura.")
    receptor = re.search(
        r"\bRECEPTOR\s*:(.*?)(?:DOMICILIO\s+FISCAL|USO\s+(?:DEL\s+)?CFDI|REGIMEN\s+FISCAL|DESCRIPCION)", t
    )
    nombre = re.search(r"\bNOMBRE\s*:\s*(.+?)\s*$", receptor[1]) if receptor else None
    if not nombre or normalizar(nombre[1]) != normalizar(cliente):
        raise RevisionNecesaria("El nombre del receptor del PDF no coincide exactamente con el cliente.")
    return {"metodo": metodos.pop(), "evidencia": " | ".join(evidencias),
            "serie": serie, "folio": str(folio)}


def analizar_pdf(ruta: Path, cliente: str, factura: str) -> dict:
    from pypdf import PdfReader
    if ruta.read_bytes()[:5] != b"%PDF-":
        raise RevisionNecesaria("La descarga no contiene un PDF.")
    lector = PdfReader(str(ruta))
    if lector.is_encrypted:
        raise RevisionNecesaria("PDF cifrado; requiere revisión.")
    # El layout mantiene juntos etiqueta y valor en el encabezado de dos columnas.
    texto = "\n".join(p.extract_text(extraction_mode="layout") or "" for p in lector.pages)
    resultado = analizar_texto(texto, cliente, factura)
    resultado.update(pdf=str(ruta.resolve()), sha256=hashlib.sha256(ruta.read_bytes()).hexdigest())
    return resultado


class Bitacora:
    """SQLite persiste ANTES del clic. Un intento incierto bloquea otro envío."""
    def __init__(self, carpeta: Path, sucursal: str):
        carpeta.mkdir(parents=True, exist_ok=True)
        self.carpeta = carpeta
        self.sucursal = sucursal
        self.db = sqlite3.connect(carpeta / "bitacora.sqlite3")
        self.db.execute("PRAGMA synchronous=FULL")
        self.db.executescript("""
            CREATE TABLE IF NOT EXISTS facturas (
                sucursal TEXT, recibo TEXT, factura TEXT, datos TEXT NOT NULL,
                PRIMARY KEY (sucursal, recibo, factura));
            CREATE TABLE IF NOT EXISTS envios (
                sucursal TEXT, recibo TEXT, estado TEXT, mensaje TEXT,
                PRIMARY KEY (sucursal, recibo));
        """)

    def registrar(self, cliente, recibo, factura, estado, **detalles):
        datos = dict(fecha=datetime.now().isoformat(timespec="seconds"), sucursal=self.sucursal,
                     cliente=cliente, recibo=recibo, factura=factura, estado=estado, **detalles)
        with self.db:
            self.db.execute("INSERT OR REPLACE INTO facturas VALUES (?, ?, ?, ?)",
                            (self.sucursal, recibo, factura, json.dumps(datos, ensure_ascii=False)))

    def envio(self, recibo):
        row = self.db.execute("SELECT estado FROM envios WHERE sucursal=? AND recibo=?",
                              (self.sucursal, recibo)).fetchone()
        return row[0] if row else None

    def iniciar_envio(self, recibo):
        # UNIQUE evita que dos procesos puedan emitir para el mismo recibo.
        with self.db:
            self.db.execute("INSERT INTO envios VALUES (?, ?, 'ENVIO_INICIADO', '')",
                            (self.sucursal, recibo))

    def confirmar_envio(self, recibo, mensaje):
        with self.db:
            self.db.execute("UPDATE envios SET estado='TIMBRADO_CONFIRMADO', mensaje=? "
                            "WHERE sucursal=? AND recibo=?", (mensaje, self.sucursal, recibo))

    def exportar(self, destino: Path):
        # Dependencia del programa distribuido; el reporte se genera en la máquina
        # del operador en cada ejecución, sin depender del entorno de Codex.
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        columnas = [("cliente", "Cliente", 36), ("factura", "Número de factura", 22),
                    ("metodo", "Método de pago", 18), ("recibo", "Recibo", 22),
                    ("estado", "Resultado", 32), ("mensaje", "Detalle", 65),
                    ("fecha", "Fecha de revisión", 23), ("sucursal", "Sucursal", 20),
                    ("pdf", "PDF consultado", 65), ("sha256", "SHA256 del PDF", 68),
                    ("evidencia", "Texto del método en PDF", 65),
                    ("cuenta_cliente", "Cuenta del cliente en ERP", 28)]
        libro = Workbook()
        hoja = libro.active
        hoja.title = "Revisión de facturas"
        hoja.append([c[1] for c in columnas])
        for (raw,) in self.db.execute("SELECT datos FROM facturas WHERE sucursal=? ORDER BY recibo,factura",
                                     (self.sucursal,)):
            datos = json.loads(raw)
            hoja.append([str(datos.get(c[0], "")) for c in columnas])
            for celda in hoja[hoja.max_row]:
                celda.data_type = "s"  # nombres que empiezan con '=' nunca son fórmulas
                celda.alignment = Alignment(vertical="top", wrap_text=True)
        for celda in hoja[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="17365D")
        for i, (_, _, ancho) in enumerate(columnas, 1):
            hoja.column_dimensions[get_column_letter(i)].width = ancho
        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = hoja.dimensions
        temporal = destino.with_suffix(".tmp.xlsx")
        libro.save(temporal)
        temporal.replace(destino)


def decision(metodos: list[str]) -> str:
    if not metodos or any(m not in {"PUE", "PPD"} for m in metodos):
        return "REVISION_NECESARIA"
    if set(metodos) == {"PUE"}:
        return "OMITIDO_PUE"
    if set(metodos) == {"PPD"}:
        return "APTO_PPD"
    # El botón emite por recibo; no conocemos una selección parcial de documentos.
    return "REVISION_RECIBO_MIXTO"
